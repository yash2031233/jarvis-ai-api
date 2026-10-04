"""Find files by what's IN them: a full-text index (SQLite FTS5) of the user's documents, refreshed in the
background, with optional re-ranking by meaning when the model provider offers embeddings.

Indexed: text, Markdown, code, CSV/JSON, Word (.docx), PDF (first pages), PowerPoint (.pptx, if python-pptx).
Folders: Documents, Desktop, Downloads (+ OneDrive's) unless `index_folders` is set.
"""

from __future__ import annotations

import logging
import os
import re
import sqlite3
import threading
import time
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    import numpy as np

from .. import config

log = logging.getLogger(__name__)
DB = config.DATA_DIR / "content_index.db"
TEXT_EXT = {".txt", ".md", ".markdown", ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".cs",
            ".go", ".rs", ".rb", ".php", ".html", ".css", ".json", ".csv", ".yaml", ".yml", ".toml", ".ini",
            ".sql", ".sh", ".ps1", ".bat", ".tex", ".rtf", ".log", ".scad"}
DOC_EXT = {".docx", ".pdf", ".pptx"}
SKIP_DIRS = {"node_modules", ".git", "__pycache__", ".venv", "venv", "AppData", "site-packages", "$RECYCLE.BIN",
             "build", "dist", ".cache", ".idea", ".vscode"}
MAX_BYTES = 8_000_000
MAX_CHARS = 60_000
STOP = set("a an the and or of to in on for with my me i is are was were be it this that what where which who "
           "about from by at as find file files doc document documents please can you".split())

_lock = threading.Lock()
_running = threading.Event()


def _conn() -> sqlite3.Connection:
    DB.parent.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(DB, check_same_thread=False)
    c.executescript("""
        CREATE TABLE IF NOT EXISTS files(path TEXT PRIMARY KEY, mtime REAL, size INTEGER, indexed REAL);
        CREATE VIRTUAL TABLE IF NOT EXISTS docs USING fts5(path UNINDEXED, name, body, tokenize='porter unicode61');
        CREATE TABLE IF NOT EXISTS vecs(path TEXT PRIMARY KEY, vec BLOB);
    """)
    return c


def folders() -> list[Path]:
    s = config.store.load()
    if s.index_folders:
        return [Path(os.path.expandvars(os.path.expanduser(p))) for p in s.index_folders]
    from ..hands.osutil import user_folders

    uf = user_folders()
    return [p for k, p in uf.items() if k in ("documents", "desktop", "downloads") and p.exists()]


def extract(p: Path) -> str:
    ext = p.suffix.lower()
    try:
        if ext in TEXT_EXT:
            return p.read_bytes()[:MAX_CHARS * 2].decode("utf-8", errors="ignore")[:MAX_CHARS]
        if ext == ".docx":
            import docx

            d = docx.Document(str(p))
            parts = [par.text for par in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)[:MAX_CHARS]
        if ext == ".pdf":
            from pypdf import PdfReader

            r = PdfReader(str(p))
            return "\n".join((pg.extract_text() or "") for pg in r.pages[:40])[:MAX_CHARS]
        if ext == ".pptx":
            from pptx import Presentation  # optional

            out = []
            for slide in Presentation(str(p)).slides:
                for sh in slide.shapes:
                    if getattr(sh, "has_text_frame", False):
                        out.append(sh.text_frame.text)
            return "\n".join(out)[:MAX_CHARS]
    except Exception as e:
        log.debug("extract failed %s: %s", p, e)
    return ""


def build(max_files: int = 40000) -> dict[str, int]:
    """Incremental: only new/changed files are (re)read; deleted files are dropped."""
    if _running.is_set():
        return {"skipped": 1}
    _running.set()
    t0 = time.time()
    added = removed = seen = 0
    try:
        with _lock:
            c = _conn()
            known = {r[0]: (r[1], r[2]) for r in c.execute("SELECT path, mtime, size FROM files")}
        present: set[str] = set()
        for root in folders():
            for dirpath, dirnames, filenames in os.walk(root):
                dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
                for n in filenames:
                    p = Path(dirpath) / n
                    ext = p.suffix.lower()
                    if ext not in TEXT_EXT and ext not in DOC_EXT:
                        continue
                    try:
                        st = p.stat()
                    except OSError:
                        continue
                    if st.st_size > MAX_BYTES:
                        continue
                    sp = str(p)
                    present.add(sp)
                    seen += 1
                    if seen > max_files:
                        break
                    if known.get(sp) == (st.st_mtime, st.st_size):
                        continue
                    body = extract(p)
                    with _lock:
                        c.execute("DELETE FROM docs WHERE path = ?", (sp,))
                        c.execute("DELETE FROM vecs WHERE path = ?", (sp,))   # changed: re-embed
                        if body.strip():
                            c.execute("INSERT INTO docs(path, name, body) VALUES (?,?,?)", (sp, p.stem, body))
                        c.execute("INSERT OR REPLACE INTO files VALUES (?,?,?,?)", (sp, st.st_mtime, st.st_size, time.time()))
                    added += 1
                    if added % 200 == 0:
                        with _lock:
                            c.commit()
        with _lock:
            for sp in set(known) - present:
                c.execute("DELETE FROM docs WHERE path = ?", (sp,))
                c.execute("DELETE FROM files WHERE path = ?", (sp,))
                c.execute("DELETE FROM vecs WHERE path = ?", (sp,))
                removed += 1
            c.commit()
            c.close()
    finally:
        _running.clear()
    log.info("content index: %d files, %d (re)indexed, %d removed in %.1fs", seen, added, removed, time.time() - t0)
    return {"files": seen, "indexed": added, "removed": removed, "seconds": round(time.time() - t0, 1)}


def _fts_query(q: str) -> str:
    words = [w for w in re.findall(r"[A-Za-z0-9]+", q.lower()) if w not in STOP and len(w) > 1]
    if not words:
        words = re.findall(r"[A-Za-z0-9]+", q.lower())[:6]
    return " OR ".join(f'"{w}"*' if len(w) > 3 else f'"{w}"' for w in words[:12])


def _embed(texts: list[str], model: str) -> "np.ndarray":
    """Embeddings from the provider's OpenAI-compatible /embeddings endpoint (sync; fine in a thread)."""
    import httpx
    import numpy as np

    s = config.store.load()
    key = config.get_api_key() or "local"
    r = httpx.post(s.base_url.rstrip("/") + "/embeddings", json={"model": model, "input": texts},
                   headers={"Authorization": f"Bearer {key}"}, timeout=120)
    r.raise_for_status()
    v = np.array([d["embedding"] for d in r.json()["data"]], dtype=np.float32)
    return v / (np.linalg.norm(v, axis=1, keepdims=True) + 1e-9)


_matrix: dict[str, Any] = {}


def embed_pending(model: str, batch: int = 32) -> int:
    """Give every indexed document a meaning vector (name + its first ~2000 characters)."""
    done = 0
    with _lock:
        c = _conn()
        todo = c.execute("SELECT d.path, d.name, substr(d.body, 1, 2000) FROM docs d "
                         "LEFT JOIN vecs v ON v.path = d.path WHERE v.path IS NULL").fetchall()
        c.close()
    for i in range(0, len(todo), batch):
        chunk = todo[i:i + batch]
        try:
            v = _embed([f"{n}\n{b}" for _, n, b in chunk], model)
        except Exception as e:
            log.info("embedding failed (keyword search still works): %s", e)
            break
        with _lock:
            c = _conn()
            c.executemany("INSERT OR REPLACE INTO vecs VALUES (?, ?)",
                          [(p, v[j].tobytes()) for j, (p, _, _) in enumerate(chunk)])
            c.commit()
            c.close()
        done += len(chunk)
    _matrix.clear()
    return done


def _vectors() -> tuple[list[str], "np.ndarray | None"]:
    import numpy as np

    if "paths" not in _matrix:
        with _lock:
            c = _conn()
            rows = c.execute("SELECT v.path, v.vec FROM vecs v JOIN files f ON f.path = v.path").fetchall()
            c.close()
        _matrix["paths"] = [r[0] for r in rows]
        _matrix["m"] = np.stack([np.frombuffer(r[1], dtype=np.float32) for r in rows]) if rows else None
    return _matrix["paths"], _matrix["m"]


async def search(query: str, limit: int = 8) -> list[dict[str, Any]]:
    """Hybrid: keyword hits (FTS5/bm25) + meaning hits (embeddings), merged and ranked together."""
    import asyncio

    import numpy as np

    fq = _fts_query(query)
    with _lock:
        c = _conn()
        rows = c.execute(
            "SELECT path, name, snippet(docs, 2, '[', ']', ' … ', 18), bm25(docs, 0.0, 4.0, 1.0) AS rank "
            "FROM docs WHERE docs MATCH ? ORDER BY rank LIMIT 30", (fq,)).fetchall() if fq else []
        c.close()
    hits: dict[str, dict[str, Any]] = {}
    if rows:
        kw = np.array([-r[3] for r in rows], dtype=np.float32)
        kw = (kw - kw.min()) / (kw.max() - kw.min() + 1e-9) if len(kw) > 1 else np.ones(1, np.float32)
        for r, k in zip(rows, kw):
            hits[r[0]] = {"path": r[0], "name": r[1], "snippet": " ".join(r[2].split()), "keyword": round(float(k), 3)}
    model = await _embed_model()
    paths, m = _vectors() if model else ([], None)
    if model and m is not None:
        try:
            q = (await asyncio.to_thread(_embed, [query], model))[0]
            sims = m @ q
            top = np.argsort(-sims)[:30]
            # keep meaning-only hits that are close to the best one (unrelated docs still score ~0.45-0.5)
            floor = max(0.5, float(sims[top[0]]) - 0.07) if len(top) else 1.0
            for i in top:
                p = paths[i]
                h = hits.get(p)
                if h is None:
                    if sims[i] < floor:
                        continue
                    h = hits[p] = {"path": p, "name": Path(p).stem, "snippet": _start(p), "keyword": 0.0}
                h["meaning"] = round(float(sims[i]), 3)
        except Exception as e:
            log.debug("meaning search unavailable: %s", e)
    for h in hits.values():
        h["score"] = round(0.6 * h.get("meaning", 0.0) + 0.4 * h["keyword"], 3) if "meaning" in h else h["keyword"]
    out = sorted(hits.values(), key=lambda h: -h["score"])[:limit]
    for h in out:
        try:
            h["modified"] = time.strftime("%Y-%m-%d", time.localtime(Path(h["path"]).stat().st_mtime))
        except OSError:
            pass
    return out


def _start(path: str) -> str:
    with _lock:
        c = _conn()
        r = c.execute("SELECT substr(body, 1, 220) FROM docs WHERE path = ?", (path,)).fetchone()
        c.close()
    return " ".join((r[0] if r else "").split())


_auto_embed: dict[str, str] = {}


async def _embed_model() -> str:
    """The configured embeddings model, or one the provider lists (e.g. nomic-embed in LM Studio)."""
    s = config.store.load()
    if s.embed_model:
        return s.embed_model
    if s.base_url in _auto_embed:
        return _auto_embed[s.base_url]
    try:
        from ..brain.client import brain

        ids = await brain.list_models()
        found = next((m for m in ids if "embed" in m.lower()), "")
    except Exception:
        found = ""
    _auto_embed[s.base_url] = found
    return found


async def _embed_model_fresh() -> str:
    """Model detection from a background thread (its own short-lived client)."""
    s = config.store.load()
    if s.embed_model:
        return s.embed_model
    try:
        import httpx

        key = config.get_api_key() or "local"
        r = httpx.get(s.base_url.rstrip("/") + "/models", headers={"Authorization": f"Bearer {key}"}, timeout=15)
        return next((m["id"] for m in r.json().get("data", []) if "embed" in m["id"].lower()), "")
    except Exception:
        return ""


def stats() -> dict[str, Any]:
    with _lock:
        c = _conn()
        n = c.execute("SELECT count(*) FROM files").fetchone()[0]
        last = c.execute("SELECT max(indexed) FROM files").fetchone()[0]
        c.close()
    return {"files": n, "last_indexed": last, "folders": [str(f) for f in folders()], "running": _running.is_set()}


def warm() -> None:
    def loop() -> None:
        time.sleep(20)  # let the app start first
        while True:
            try:
                build()
                import asyncio

                model = asyncio.run(_embed_model_fresh())
                if model:
                    embed_pending(model)
            except Exception as e:
                log.info("content index failed: %s", e)
            time.sleep(1800)

    threading.Thread(target=loop, daemon=True, name="content-index").start()
